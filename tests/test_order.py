from io import BytesIO
import tempfile
import unittest
from datetime import date
from pathlib import Path
from docx import Document
from fastapi import HTTPException
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from src.db.base import Base
from src.models.domain import Assignment, Group, Organization, Practice, Student
from src.repositories.assignments import AssignmentRepository
from src.repositories.reports import ReportRepository
from src.services.documents import render_order


class OrderGenerationTests(unittest.IsolatedAsyncioTestCase):
    async def test_order_has_all_students_in_import_order_and_assignment_data(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "order.db"
            engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
            try:
                async with engine.begin() as connection:
                    await connection.run_sync(Base.metadata.create_all)

                sessions = async_sessionmaker(engine, expire_on_commit=False)
                async with sessions() as session:
                    groups = [
                        Group(name="ГР-25", year=2025),
                        Group(name="ГР-24", year=2024),
                    ]
                    session.add_all(groups)
                    await session.flush()
                    students = [
                        Student(
                            full_name="Иванов Иван Иванович",
                            group_id=groups[0].id,
                            record_book=1001,
                        ),
                        Student(
                            full_name="Петров Петр Петрович",
                            group_id=groups[0].id,
                            record_book=1002,
                        ),
                        Student(
                            full_name="Сидорова Анна Сергеевна",
                            group_id=groups[1].id,
                            record_book=1003,
                        ),
                    ]
                    organizations = [
                        Organization(name="ООО Ромашка", address="г. Уфа"),
                        Organization(name="АО Пример", address=None),
                    ]
                    practice = Practice(
                        type="Учебная",
                        start_date=date(2026, 6, 1),
                        end_date=date(2026, 6, 30),
                        group_id=groups[0].id,
                    )
                    session.add_all([*students, *organizations, practice])
                    await session.flush()
                    session.add_all(
                        [
                            Assignment(
                                student_id=students[0].id,
                                organization_id=organizations[0].id,
                                practice_id=practice.id,
                            ),
                            Assignment(
                                student_id=students[1].id,
                                organization_id=organizations[1].id,
                                practice_id=practice.id,
                            ),
                        ]
                    )
                    await session.commit()

                    rows = await AssignmentRepository(session).get_order_rows()

                self.assertEqual(
                    [row["full_name"] for row in rows],
                    [student.full_name for student in students],
                )
                self.assertEqual(rows[0]["org_display"], "ООО Ромашка, г. Уфа")
                self.assertEqual(rows[0]["practice_type"], "Учебная")
                self.assertEqual(rows[0]["start_date"], "01.06.2026")
                self.assertEqual(rows[1]["org_display"], "АО Пример")
                self.assertEqual(rows[2]["org_display"], "")

                document_bytes = render_order(rows)
                output = Path(temp_dir) / "generated-order.docx"
                output.write_bytes(document_bytes)
                document = Document(output)
                table = document.tables[0]
                generated_names = [
                    table.cell(i, 1).text for i in range(1, len(table.rows))
                ]
                self.assertEqual(
                    generated_names,
                    [student.full_name for student in students],
                )
                self.assertEqual(table.cell(1, 0).text, "1")
                self.assertEqual(table.cell(2, 0).text, "2")
                self.assertEqual(table.cell(3, 0).text, "3")
                self.assertEqual(table.cell(1, 2).text, "ООО Ромашка, г. Уфа")
                self.assertEqual(table.cell(3, 2).text, "")
            finally:
                await engine.dispose()

    async def test_directions_export_returns_single_docx_with_page_breaks(self) -> None:
        from src.routers.documents import download_directions

        with tempfile.TemporaryDirectory() as temp_dir:
            engine = create_async_engine(
                f"sqlite+aiosqlite:///{Path(temp_dir) / 'directions.db'}"
            )
            try:
                async with engine.begin() as connection:
                    await connection.run_sync(Base.metadata.create_all)
                sessions = async_sessionmaker(engine, expire_on_commit=False)
                async with sessions() as session:
                    group = Group(name="ГР-25", year=2025)
                    organization = Organization(name="ООО Ромашка", address="г. Уфа")
                    session.add_all([group, organization])
                    await session.flush()
                    students = [
                        Student(
                            full_name="Иванов Иван Иванович",
                            group_id=group.id,
                            record_book=3001,
                        ),
                        Student(
                            full_name="Петров Петр Петрович",
                            group_id=group.id,
                            record_book=3002,
                        ),
                    ]
                    practice = Practice(
                        type="Учебная",
                        start_date=date(2026, 6, 1),
                        end_date=date(2026, 6, 30),
                        group_id=group.id,
                    )
                    session.add_all([*students, practice])
                    await session.flush()
                    session.add_all(
                        [
                            Assignment(
                                student_id=student.id,
                                organization_id=organization.id,
                                practice_id=practice.id,
                            )
                            for student in students
                        ]
                    )
                    await session.commit()
                    directions = await AssignmentRepository(
                        session
                    ).get_directions_data("ГР-25")
                    response = await download_directions("ГР-25", session)

                self.assertEqual(len(directions), 2)
                self.assertEqual(directions[0]["group_name"], "ГР-25")
                self.assertEqual(directions[0]["practice_start_date"], "01.06.2026")
                self.assertEqual(directions[0]["practice_end_date"], "30.06.2026")
                self.assertEqual(
                    response.media_type,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
                self.assertEqual(
                    response.headers["content-disposition"],
                    'attachment; filename="directions.docx"',
                )
                document = Document(BytesIO(response.body))
                self.assertEqual(len(document.tables), 2)
                document_text = "\n".join(
                    cell.text
                    for table in document.tables
                    for row in table.rows
                    for cell in row.cells
                )
                self.assertIn("Иванов Иван Иванович", document_text)
                self.assertIn("Петров Петр Петрович", document_text)
                self.assertEqual(
                    len(document._element.body.xpath(".//w:br[@w:type='page']")),
                    1,
                )
            finally:
                await engine.dispose()

    async def test_report_counts_unique_paid_students(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            engine = create_async_engine(
                f"sqlite+aiosqlite:///{Path(temp_dir) / 'report.db'}"
            )
            try:
                async with engine.begin() as connection:
                    await connection.run_sync(Base.metadata.create_all)
                sessions = async_sessionmaker(engine, expire_on_commit=False)
                async with sessions() as session:
                    group = Group(name="ГР-25", year=2025)
                    organization = Organization(name="ООО Ромашка")
                    session.add_all([group, organization])
                    await session.flush()
                    student = Student(
                        full_name="Иванов Иван Иванович",
                        group_id=group.id,
                        record_book=4001,
                    )
                    practice = Practice(
                        type="Учебная",
                        start_date=date(2026, 6, 1),
                        end_date=date(2026, 6, 30),
                        group_id=group.id,
                    )
                    session.add_all([student, practice])
                    await session.flush()
                    session.add_all(
                        [
                            Assignment(
                                student_id=student.id,
                                organization_id=organization.id,
                                practice_id=practice.id,
                                payment_type="платное",
                            ),
                            Assignment(
                                student_id=student.id,
                                organization_id=organization.id,
                                practice_id=practice.id,
                                payment_type="платное",
                            ),
                        ]
                    )
                    await session.commit()
                    report = await ReportRepository(session).get_groups_report_data()

                self.assertEqual(report[0]["count_stud"], 1)
                self.assertEqual(report[0]["payed_students"], 1)
            finally:
                await engine.dispose()

    async def test_order_endpoint_returns_downloadable_docx(self) -> None:
        from src.routers.documents import download_order

        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "order-endpoint.db"
            engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
            try:
                async with engine.begin() as connection:
                    await connection.run_sync(Base.metadata.create_all)
                sessions = async_sessionmaker(engine, expire_on_commit=False)
                async with sessions() as session:
                    group = Group(name="ГР-25", year=2025)
                    session.add(group)
                    await session.flush()
                    session.add(
                        Student(
                            full_name="Тест Тест Тестович",
                            group_id=group.id,
                            record_book=2001,
                        )
                    )
                    await session.commit()
                    response = await download_order(session)

                self.assertEqual(
                    response.media_type,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
                self.assertEqual(
                    response.headers["content-disposition"],
                    'attachment; filename="order.docx"',
                )
                document = Document(BytesIO(response.body))
                self.assertEqual(document.tables[0].cell(1, 1).text, "Тест Тест Тестович")
                self.assertTrue(response.body.startswith(b"PK"))
            finally:
                await engine.dispose()

    async def test_order_endpoint_returns_404_when_there_are_no_students(self) -> None:
        from src.routers.documents import download_order

        with tempfile.TemporaryDirectory() as temp_dir:
            database_path = Path(temp_dir) / "empty-order.db"
            engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
            try:
                async with engine.begin() as connection:
                    await connection.run_sync(Base.metadata.create_all)
                sessions = async_sessionmaker(engine, expire_on_commit=False)
                async with sessions() as session:
                    with self.assertRaises(HTTPException) as error:
                        await download_order(session)

                self.assertEqual(error.exception.status_code, 404)
            finally:
                await engine.dispose()


if __name__ == "__main__":
    unittest.main()